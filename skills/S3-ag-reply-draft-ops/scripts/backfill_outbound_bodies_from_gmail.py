#!/usr/bin/env python3
from __future__ import annotations

import argparse
import base64
import json
import re
from pathlib import Path

import pandas as pd
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from urllib import request


DEFAULT_TOKEN_FILE = "${ROCKBASE_HOME}/Documents/GitHub/AG-Skills-Hub/02 💼 Office/ag-Google-Suite/auth/<YOUR_ACCOUNT_EMAIL>"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Backfill outbound Mail1/Mail2/Mail3 sent bodies from Gmail sent message ids.")
    parser.add_argument("--source-csv", required=True)
    parser.add_argument("--review-xlsx", default="")
    parser.add_argument("--stage", required=True)
    parser.add_argument("--outbound-wave", default="mail1")
    parser.add_argument("--target-field", default="Mail1_Body_Final")
    parser.add_argument("--token-file", default=DEFAULT_TOKEN_FILE)
    parser.add_argument("--summary-md", required=True)
    return parser.parse_args()


def normalize_text(value: object) -> str:
    return str(value or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def build_creds(token_file: str) -> Credentials:
    creds = Credentials.from_authorized_user_file(
        token_file,
        scopes=["https://www.googleapis.com/auth/gmail.readonly"],
    )
    if not creds.valid:
        if creds.expired and creds.refresh_token:
            creds.refresh(Request())
            Path(token_file).write_text(creds.to_json(), encoding="utf-8")
        else:
            raise RuntimeError(f"Token file is invalid and cannot refresh: {token_file}")
    return creds


def gmail_request(access_token: str, path: str, timeout: int = 30) -> dict:
    url = f"https://gmail.googleapis.com/gmail/v1/users/me/{path.lstrip('/')}"
    req = request.Request(url, headers={"Authorization": f"Bearer {access_token}"})
    with request.urlopen(req, timeout=timeout) as resp:
        payload = resp.read()
    return json.loads(payload.decode("utf-8")) if payload else {}


def extract_text_from_payload(payload: dict) -> str:
    mime = payload.get("mimeType", "")
    body = payload.get("body", {}) or {}
    data = body.get("data")
    parts = payload.get("parts", []) or []

    if mime == "text/plain" and data:
        return base64.urlsafe_b64decode(data + "===").decode("utf-8", errors="ignore").strip()
    if mime == "text/html" and data:
        html = base64.urlsafe_b64decode(data + "===").decode("utf-8", errors="ignore")
        text = re.sub(r"<br\s*/?>", "\n", html, flags=re.I)
        text = re.sub(r"</p\s*>", "\n\n", text, flags=re.I)
        text = re.sub(r"<[^>]+>", " ", text)
        text = re.sub(r"\n{3,}", "\n\n", text)
        text = re.sub(r"[ \t]+", " ", text)
        return text.strip()
    for part in parts:
        text = extract_text_from_payload(part)
        if text:
            return text
    return ""


def extract_outbound_message_id(outbound_message_ids: str, wave: str) -> str:
    wave = wave.lower().strip()
    for chunk in (outbound_message_ids or "").split("|"):
        if ":" not in chunk:
            continue
        key, value = chunk.split(":", 1)
        if key.strip().lower() == wave:
            return value.strip()
    return ""


def ensure_xlsx_column(ws, field_name: str, after_field: str = "账号链接") -> int:
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    if field_name in headers:
        return headers.index(field_name) + 1

    insert_at = headers.index(after_field) + 2 if after_field in headers else ws.max_column + 1
    ws.insert_cols(insert_at)
    ws.cell(1, insert_at).value = field_name
    cell = ws.cell(1, insert_at)
    cell.fill = PatternFill(fill_type="solid", fgColor="D9EAF7")
    cell.font = Font(bold=True)
    cell.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.column_dimensions[ws.cell(1, insert_at).column_letter].width = 56
    return insert_at


def main() -> None:
    args = parse_args()
    creds = build_creds(args.token_file)
    access_token = creds.token
    if not access_token:
        raise RuntimeError("Could not obtain access token from token file.")

    csv_path = Path(args.source_csv)
    df = pd.read_csv(csv_path, dtype=str, keep_default_na=False)
    mask = df["Pipeline_Stage"].astype(str).eq(args.stage)
    stage_df = df[mask].copy()

    fetched_cache: dict[str, str] = {}
    updated_rows: list[dict[str, str]] = []
    for idx, row in df[mask].iterrows():
        outbound_id = extract_outbound_message_id(row.get("Outbound_Message_IDs", ""), args.outbound_wave)
        if not outbound_id:
            continue
        if outbound_id not in fetched_cache:
            payload = gmail_request(access_token, f"messages/{outbound_id}?format=full")
            fetched_cache[outbound_id] = extract_text_from_payload(payload.get("payload", {}) or {})
        body_text = fetched_cache[outbound_id]
        if not body_text:
            continue
        df.at[idx, args.target_field] = body_text
        updated_rows.append(
            {
                "账号ID": normalize_text(row.get("账号ID", "")),
                "频道/作者名称": normalize_text(row.get("频道/作者名称", "")),
                "outbound_message_id": outbound_id,
            }
        )

    df.to_csv(csv_path, index=False, encoding="utf-8-sig")

    if args.review_xlsx:
        xlsx_path = Path(args.review_xlsx)
        wb = load_workbook(xlsx_path)
        ws = wb[wb.sheetnames[0]]
        target_col = ensure_xlsx_column(ws, args.target_field, after_field="账号链接")
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        idx_map = {name: i + 1 for i, name in enumerate(headers)}
        csv_lookup = {
            normalize_text(row["账号ID"]): normalize_text(row[args.target_field])
            for _, row in df[mask].iterrows()
            if normalize_text(row.get("账号ID", ""))
        }
        for row_no in range(2, ws.max_row + 1):
            account_id = normalize_text(ws.cell(row_no, idx_map["账号ID"]).value)
            if account_id not in csv_lookup:
                continue
            ws.cell(row_no, target_col).value = csv_lookup[account_id]
            ws.cell(row_no, target_col).alignment = Alignment(vertical="top", wrap_text=True)
        wb.save(xlsx_path)

    summary_lines = [
        "---",
        "tags:",
        "  - replyops",
        "  - outbound-body",
        "  - gmail",
        "date: 2026-03-20",
        "status: done",
        "---",
        "",
        f"# 2026-03-20 Backfill {args.target_field}",
        "",
        f"- Stage: `{args.stage}`",
        f"- Wave source: `{args.outbound_wave}`",
        f"- Target CSV field: `{args.target_field}`",
        f"- Updated rows: {len(updated_rows)}",
        "",
        "## Rows",
        "",
    ]
    for row in updated_rows:
        summary_lines.append(f"- {row['账号ID']} | {row['频道/作者名称']} | `{row['outbound_message_id']}`")
    Path(args.summary_md).write_text("\n".join(summary_lines) + "\n", encoding="utf-8")

    print(json.dumps({"updated_rows": len(updated_rows), "summary_md": args.summary_md}, ensure_ascii=False))


if __name__ == "__main__":
    main()
